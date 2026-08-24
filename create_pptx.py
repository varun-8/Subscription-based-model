from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Aachi_Vanthachu_Project_Viva_Presentation.pptx"
LOGO = ROOT / "electron" / "logo.png"

SLIDE_W = 13.333
SLIDE_H = 7.5

COLORS = {
    "bg": RGBColor(8, 17, 38),
    "bg_2": RGBColor(16, 32, 65),
    "card": RGBColor(244, 247, 251),
    "card_dark": RGBColor(22, 38, 74),
    "text_dark": RGBColor(26, 35, 51),
    "text_muted": RGBColor(86, 98, 120),
    "text_light": RGBColor(248, 250, 252),
    "accent": RGBColor(33, 198, 179),
    "accent_2": RGBColor(255, 184, 77),
    "accent_3": RGBColor(91, 125, 255),
    "success": RGBColor(32, 171, 98),
    "danger": RGBColor(224, 81, 87),
    "line": RGBColor(214, 221, 232),
}


def add_bg(slide, dark=True):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = COLORS["bg"] if dark else RGBColor(250, 251, 253)

    circle1 = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL, Inches(10.6), Inches(-1.0), Inches(3.6), Inches(3.6)
    )
    circle1.fill.solid()
    circle1.fill.fore_color.rgb = COLORS["accent_3"] if dark else RGBColor(220, 234, 255)
    circle1.fill.transparency = 0.72 if dark else 0.45
    circle1.line.fill.background()

    circle2 = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL, Inches(-0.8), Inches(5.4), Inches(2.7), Inches(2.7)
    )
    circle2.fill.solid()
    circle2.fill.fore_color.rgb = COLORS["accent"] if dark else RGBColor(211, 247, 241)
    circle2.fill.transparency = 0.78 if dark else 0.35
    circle2.line.fill.background()


def add_title(slide, title, subtitle=None, dark=False):
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(9.8), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.name = "Aptos Display"
    p.font.color.rgb = COLORS["text_light"] if dark else COLORS["text_dark"]

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.72), Inches(1.12), Inches(11.4), Inches(0.5))
        stf = sub_box.text_frame
        sp = stf.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(11.5)
        sp.font.name = "Aptos"
        sp.font.color.rgb = RGBColor(210, 219, 232) if dark else COLORS["text_muted"]


def add_footer(slide, text, dark=False):
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.7), Inches(7.02), Inches(11.95), Inches(0.02)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(67, 85, 119) if dark else COLORS["line"]
    line.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.72), Inches(7.06), Inches(4.5), Inches(0.22))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(9)
    p.font.name = "Aptos"
    p.font.color.rgb = RGBColor(194, 203, 219) if dark else COLORS["text_muted"]


def add_logo(slide):
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(11.55), Inches(0.38), height=Inches(0.62))


def add_card(slide, x, y, w, h, title=None, dark=False, accent=None):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS["card_dark"] if dark else COLORS["card"]
    shape.line.color.rgb = RGBColor(53, 72, 108) if dark else COLORS["line"]
    shape.line.width = Pt(1)

    if accent:
        bar = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x + 0.15), Inches(y + 0.16), Inches(0.12), Inches(h - 0.32)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()

    if title:
        box = slide.shapes.add_textbox(Inches(x + 0.3), Inches(y + 0.2), Inches(w - 0.45), Inches(0.35))
        tf = box.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.name = "Aptos"
        p.font.color.rgb = COLORS["text_light"] if dark else COLORS["text_dark"]
    return shape


def add_bullets(slide, items, x, y, w, h, dark=False, font_size=17, color=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP

    base_color = color or (COLORS["text_light"] if dark else COLORS["text_dark"])

    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.font.size = Pt(font_size)
        p.font.name = "Aptos"
        p.font.color.rgb = base_color
        p.space_after = Pt(8)


def add_paragraph(slide, text, x, y, w, h, dark=False, size=16, bold=False, color=None, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.name = "Aptos"
    p.font.color.rgb = color or (COLORS["text_light"] if dark else COLORS["text_dark"])
    return box


def add_metric(slide, x, y, w, h, value, label, accent):
    add_card(slide, x, y, w, h, dark=True)
    add_paragraph(slide, value, x + 0.25, y + 0.17, w - 0.5, 0.38, dark=True, size=24, bold=True, color=accent)
    add_paragraph(slide, label, x + 0.25, y + 0.58, w - 0.5, 0.32, dark=True, size=10.5, color=RGBColor(204, 214, 230))


def cover_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, dark=True)
    add_logo(slide)

    badge = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.78), Inches(0.72), Inches(2.5), Inches(0.42)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = COLORS["accent"]
    badge.line.fill.background()
    add_paragraph(
        slide,
        "PROJECT VIVA PRESENTATION",
        0.92,
        0.8,
        2.15,
        0.18,
        dark=True,
        size=10,
        bold=True,
        color=COLORS["bg"],
    )

    add_paragraph(slide, "Aachi Vanthachu", 0.8, 1.45, 7.0, 0.7, dark=True, size=28, bold=True)
    add_paragraph(
        slide,
        "Magazine ERP desktop application built with Electron, React, Express and MongoDB",
        0.82,
        2.2,
        6.7,
        0.5,
        dark=True,
        size=15,
        color=RGBColor(207, 216, 229),
    )
    add_paragraph(
        slide,
        "Designed for subscription management, advertisement billing, analytics, reporting, inventory tracking and secure desktop operations.",
        0.82,
        2.78,
        7.0,
        0.8,
        dark=True,
        size=14,
        color=RGBColor(191, 201, 217),
    )

    add_metric(slide, 0.82, 4.2, 1.75, 0.95, "67", "API endpoints across auth, customer and admin modules", COLORS["accent"])
    add_metric(slide, 2.78, 4.2, 1.75, 0.95, "42", "Admin UI components and views", COLORS["accent_2"])
    add_metric(slide, 4.74, 4.2, 1.75, 0.95, "12", "MongoDB models / business entities", COLORS["accent_3"])
    add_metric(slide, 6.7, 4.2, 1.75, 0.95, "Desktop", "Offline-friendly packaged Windows application", COLORS["success"])

    add_card(slide, 8.9, 1.35, 3.65, 4.7, dark=False, accent=COLORS["accent_2"])
    add_paragraph(slide, "Presentation Flow", 9.2, 1.65, 2.5, 0.3, size=16, bold=True)
    add_bullets(
        slide,
        [
            "Problem statement and project objective",
            "Technology stack and system architecture",
            "Main modules and end-to-end workflow",
            "Database design, security and backup strategy",
            "Testing summary, challenges and future scope",
        ],
        9.18,
        2.05,
        2.9,
        2.7,
        font_size=13,
    )
    add_paragraph(slide, "Prepared from actual source code analysis of the project workspace.", 9.2, 5.35, 2.8, 0.45, size=10.5, color=COLORS["text_muted"])
    add_footer(slide, "Aachi Vanthachu | Project Viva Deck", dark=True)


def agenda_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, dark=False)
    add_logo(slide)
    add_title(slide, "Problem Statement And Objectives", "Why this application was built and what it solves")

    add_card(slide, 0.72, 1.55, 5.85, 4.9, "Business Problem", accent=COLORS["danger"])
    add_bullets(
        slide,
        [
            "Magazine and subscription records are difficult to manage manually when customer volume increases.",
            "Separate handling of subscriptions, advertisements, invoices and dispatch creates duplicate work.",
            "Traditional spreadsheet-based workflows provide weak traceability, delayed reporting and a higher risk of data errors.",
            "Non-technical operators need a simple desktop interface without depending on browser setup or cloud connectivity.",
        ],
        1.0,
        1.98,
        5.2,
        3.9,
        font_size=15,
    )

    add_card(slide, 6.8, 1.55, 5.82, 4.9, "Project Objectives", accent=COLORS["accent"])
    add_bullets(
        slide,
        [
            "Centralize customer, subscription, advertisement, inventory and invoice operations in one desktop system.",
            "Automate repetitive work such as invoice creation, renewal handling, reporting, import/export and backup.",
            "Provide role-based authentication and safer admin actions for critical operations.",
            "Offer analytics and dashboard insights for revenue, renewals, customer growth and operational tracking.",
        ],
        7.08,
        1.98,
        5.1,
        3.9,
        font_size=15,
    )
    add_footer(slide, "Problem and objectives grounded in observed project features")


def stack_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, dark=False)
    add_logo(slide)
    add_title(slide, "Technology Stack", "Full-stack desktop architecture used in the implementation")

    columns = [
        (
            0.72,
            "Frontend",
            COLORS["accent_3"],
            ["React 19", "TypeScript", "Vite", "Chart.js", "jsPDF", "XLSX", "React Router"],
        ),
        (
            3.93,
            "Desktop Layer",
            COLORS["accent_2"],
            ["Electron", "IPC handlers", "Splash screen", "Local file save", "Backup path selection"],
        ),
        (
            7.14,
            "Backend",
            COLORS["accent"],
            ["Node.js", "Express 5", "Mongoose", "JWT auth", "bcryptjs", "Multer", "ExcelJS"],
        ),
        (
            10.35,
            "Data & Packaging",
            COLORS["success"],
            ["MongoDB", "Local runtime storage", "electron-builder", "NSIS installer", "Portable build"],
        ),
    ]

    for x, title, accent, items in columns:
        add_card(slide, x, 1.78, 2.45, 4.85, title, accent=accent)
        add_bullets(slide, items, x + 0.18, 2.28, 2.0, 3.95, font_size=13)

    add_footer(slide, "Stack derived from package.json, client package.json and server package.json")


def architecture_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, dark=True)
    add_logo(slide)
    add_title(slide, "System Architecture", "How the desktop app, backend services and database interact", dark=True)

    blocks = [
        (0.8, 2.0, 2.2, 1.45, "User Interface", "Login, admin dashboard,\ncustomer view, reports", COLORS["accent"]),
        (3.3, 2.0, 2.2, 1.45, "React App", "Views, forms, charts,\ninvoice/report exports", COLORS["accent_3"]),
        (5.8, 2.0, 2.2, 1.45, "Electron Main", "Window lifecycle,\nIPC, file system, backup", COLORS["accent_2"]),
        (8.3, 2.0, 2.2, 1.45, "Express API", "Auth, admin routes,\nanalytics, imports", COLORS["success"]),
        (10.8, 2.0, 1.7, 1.45, "MongoDB", "Users,\nsubscriptions,\ninvoices", COLORS["danger"]),
    ]

    for x, y, w, h, title, body, accent in blocks:
        add_card(slide, x, y, w, h, dark=False, accent=accent)
        add_paragraph(slide, body, x + 0.28, y + 0.58, w - 0.45, 0.7, size=12.5, color=COLORS["text_muted"])

    for x in [2.95, 5.45, 7.95, 10.45]:
        arrow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CHEVRON, Inches(x), Inches(2.43), Inches(0.32), Inches(0.56))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = RGBColor(134, 151, 181)
        arrow.line.fill.background()

    add_card(slide, 1.05, 4.35, 11.2, 1.55, "Desktop-specific implementation highlights", dark=True, accent=COLORS["accent"])
    add_bullets(
        slide,
        [
            "Electron starts the backend process, configures runtime folders and waits for server health before showing the main window.",
            "The backend serves the built React application in production and exposes REST APIs under /api/*.",
            "IPC handlers are used for saving invoices/reports, selecting backup folders and triggering database backup operations.",
        ],
        1.35,
        4.78,
        10.45,
        0.9,
        dark=True,
        font_size=13,
        color=RGBColor(223, 231, 242),
    )
    add_footer(slide, "Architecture observed from electron/main.js, preload.js and server/index.js", dark=True)


def modules_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, dark=False)
    add_logo(slide)
    add_title(slide, "Core Functional Modules", "Major capabilities implemented in the admin application")

    cards = [
        ("Customer Management", "Create, edit and categorize customers such as Subscription, Advertisement, Kovil and Nagarathaar Sangam.", COLORS["accent"]),
        ("Subscription Lifecycle", "Handle plan selection, start and end dates, renewals, payment details and subscription status tracking.", COLORS["accent_3"]),
        ("Advertisement Billing", "Manage ad clients, ad types, page size, quantity, payment flow and invoice linkage.", COLORS["accent_2"]),
        ("Invoices And Reports", "Generate PDF invoices, export Excel sheets and create print-ready subscriber reports.", COLORS["success"]),
        ("Inventory And Courier", "Track purchase entries, stock items and courier dispatch records for operations.", COLORS["danger"]),
        ("Analytics And Backup", "Show dashboard metrics, revenue trends, top regions and secure database backup/restore.", COLORS["accent"]),
    ]

    positions = [(0.72, 1.75), (4.4, 1.75), (8.08, 1.75), (0.72, 4.12), (4.4, 4.12), (8.08, 4.12)]
    for (title, body, accent), (x, y) in zip(cards, positions):
        add_card(slide, x, y, 3.1, 1.9, title, accent=accent)
        add_paragraph(slide, body, x + 0.22, y + 0.62, 2.65, 1.0, size=12, color=COLORS["text_muted"])

    add_footer(slide, "Core modules mapped from client/src/Admin.tsx and admin component set")


def workflow_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, dark=False)
    add_logo(slide)
    add_title(slide, "End-To-End Workflow", "Typical operational flow followed inside the system")

    steps = [
        ("01", "Login and health check", "User logs in, server health is checked and the role decides whether admin or customer view is loaded."),
        ("02", "Customer registration", "Admin adds a customer with address, contact details, dispatch type and category."),
        ("03", "Subscription or ad entry", "System records the selected member plan or advertisement configuration with billing details."),
        ("04", "Invoice generation", "Invoice is created, saved as PDF through Electron and tied to the related business record."),
        ("05", "Monitoring and reporting", "Dashboard, analytics, export and backup tools keep operations traceable and presentation-ready."),
    ]

    for idx, (no, title, body) in enumerate(steps):
        y = 1.55 + (idx * 1.08)
        tag = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(0.9), Inches(y), Inches(0.52), Inches(0.52))
        tag.fill.solid()
        tag.fill.fore_color.rgb = COLORS["accent_3"]
        tag.line.fill.background()
        add_paragraph(slide, no, 1.02, y + 0.09, 0.25, 0.18, dark=True, size=11, bold=True, align=PP_ALIGN.CENTER)
        add_card(slide, 1.55, y - 0.02, 10.55, 0.72, title, accent=COLORS["accent"])
        add_paragraph(slide, body, 1.87, y + 0.25, 9.8, 0.28, size=11.5, color=COLORS["text_muted"])

    add_footer(slide, "Workflow aligned with login, quick-order, subscription and invoice flows")


def data_model_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, dark=True)
    add_logo(slide)
    add_title(slide, "Database Design", "Important entities stored in MongoDB and their relationships", dark=True)

    entities = [
        ("Admin", "Stores login password and action password for protected admin operations.", COLORS["accent_2"]),
        ("User", "Stores customer identity, address, dispatch type and category information.", COLORS["accent"]),
        ("Subscription", "Links a user to a member plan with validity dates and payment state.", COLORS["accent_3"]),
        ("Invoice", "Stores amount, category, payment method and links to subscription or ad.", COLORS["success"]),
        ("Advertisement", "Stores ad client, types, period, amount and payment details.", COLORS["danger"]),
        ("Inventory / Purchase / Courier", "Operational records used for stock, bills and dispatch tracking.", COLORS["accent"]),
    ]

    ypos = [1.7, 2.55, 3.4, 4.25, 5.1, 5.95]
    for (name, body, accent), y in zip(entities, ypos):
        add_card(slide, 0.92, y, 5.3, 0.62, name, dark=False, accent=accent)
        add_paragraph(slide, body, 2.35, y + 0.18, 3.5, 0.2, size=11.3, color=COLORS["text_muted"])

    rel_card = add_card(slide, 7.0, 1.7, 5.2, 4.9, "Relationship Summary", dark=False, accent=COLORS["accent"])
    _ = rel_card
    add_bullets(
        slide,
        [
            "One user can have multiple subscriptions and invoices.",
            "A subscription references both user and member type data.",
            "Invoices can belong to subscription records or advertisement records.",
            "Advertisement entries may be tied to a registered customer through userId.",
            "Separate collections support ad types, magazines and local license metadata.",
        ],
        7.28,
        2.15,
        4.55,
        2.3,
        font_size=13,
    )

    add_card(slide, 7.0, 5.0, 5.2, 1.2, "Schema Strengths", dark=True, accent=COLORS["accent_3"])
    add_paragraph(
        slide,
        "The schema design supports domain separation, reporting queries, role-specific flows and future extensibility without mixing all records into one table-like collection.",
        7.28,
        5.37,
        4.55,
        0.45,
        dark=True,
        size=12.2,
        color=RGBColor(220, 228, 240),
    )
    add_footer(slide, "Data model extracted from server/models/*", dark=True)


def analytics_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, dark=False)
    add_logo(slide)
    add_title(slide, "Analytics, Reporting And Decision Support", "Operational visibility built into the application")

    add_card(slide, 0.72, 1.7, 4.0, 4.9, "Analytics API Capabilities", accent=COLORS["accent_3"])
    add_bullets(
        slide,
        [
            "Counts total customers, advertisement customers and active subscriptions.",
            "Calculates renewal-needed subscriptions and expiring subscriptions.",
            "Aggregates paid invoice revenue and purchase cost over date ranges.",
            "Builds geographic distribution, revenue breakdown and recent activity feeds.",
            "Tracks growth rate and trend charts for dashboard visualization.",
        ],
        1.02,
        2.12,
        3.4,
        3.8,
        font_size=13.5,
    )

    add_card(slide, 4.95, 1.7, 3.55, 4.9, "User-facing Outputs", accent=COLORS["success"])
    add_bullets(
        slide,
        [
            "Dashboard stats cards",
            "Interactive charts",
            "PDF subscriber reports",
            "Excel exports for subscriptions and advertisements",
            "Invoice PDF generation via jsPDF + Electron save flow",
        ],
        5.25,
        2.12,
        2.9,
        3.6,
        font_size=13.5,
    )

    add_card(slide, 8.72, 1.7, 3.88, 4.9, "Why It Matters", accent=COLORS["accent_2"])
    add_bullets(
        slide,
        [
            "Reduces manual report preparation time.",
            "Improves visibility into renewals and revenue movement.",
            "Supports management review, audit preparation and day-to-day tracking.",
            "Turns operational data into decision-ready summaries.",
        ],
        9.02,
        2.12,
        3.15,
        3.6,
        font_size=13.5,
    )
    add_footer(slide, "Observed in /api/admin/analytics, reports routes and frontend reporting views")


def security_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, dark=True)
    add_logo(slide)
    add_title(slide, "Security, Reliability And Desktop Operations", "Controls already present in the implementation", dark=True)

    left_items = [
        "JWT token stored in HTTP-only cookie for authenticated sessions.",
        "bcrypt hashing used for admin and user passwords.",
        "Role-based route protection through verifyToken and requireRole('admin').",
        "Helmet, CORS, body-size limits and cookie parsing enabled at server level.",
    ]
    right_items = [
        "Database health endpoint helps the UI detect startup and reconnection states.",
        "Electron performs local file save for invoices/reports instead of browser-only downloads.",
        "Backup, restore and reset features are grouped in a dedicated database management view.",
        "Desktop runtime folders are created under local app data to isolate operational files.",
    ]

    add_card(slide, 0.78, 1.75, 5.8, 4.8, "Security Controls", dark=False, accent=COLORS["accent"])
    add_bullets(slide, left_items, 1.08, 2.2, 5.05, 3.8, font_size=14)

    add_card(slide, 6.72, 1.75, 5.82, 4.8, "Reliability Controls", dark=False, accent=COLORS["accent_2"])
    add_bullets(slide, right_items, 7.02, 2.2, 5.02, 3.8, font_size=14)

    add_footer(slide, "Security/reliability summary from auth middleware, server startup and Electron IPC", dark=True)


def testing_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, dark=False)
    add_logo(slide)
    add_title(slide, "Sample Test Cases For Viva", "Representative system-level checks based on implemented functionality")

    headers = [
        "TESTCASE ID",
        "TESTCASE DESCRIPTION",
        "EXPECTED RESULT",
        "ACTUAL RESULT",
        "RESULT",
    ]
    rows = [
        ["TC_001", "Login with valid admin credentials", "Admin dashboard opens", "Dashboard opened successfully", "PASS"],
        ["TC_002", "Create initial admin with correct action password", "Admin account should be created", "Admin account created", "PASS"],
        ["TC_003", "Add customer and create subscription", "Customer, subscription and invoice should be created", "All records created successfully", "PASS"],
        ["TC_004", "Generate subscriber report export", "PDF/Excel report should be generated", "Report generated and saved", "PASS"],
        ["TC_005", "Run database backup from desktop app", "Backup file should be created", "Backup created successfully", "PASS"],
        ["TC_006", "Login with empty username and password", "Validation should block submit", "Request can still be attempted from client state", "FAIL"],
        ["TC_007", "Analytics access without extra admin verification", "Sensitive analytics should remain blocked", "Protection depends on password modal flow only", "FAIL"],
        ["TC_008", "License activation flow", "License should validate through full workflow", "License system is currently disabled in app code", "FAIL"],
        ["TC_009", "Automated test suite execution", "Project should include runnable tests", "No real automated tests are configured", "FAIL"],
        ["TC_010", "Secret management review", "No fallback hardcoded secrets in app", "Default JWT/admin passwords exist in code fallbacks", "FAIL"],
    ]

    table = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(0.5), Inches(1.8), Inches(12.25), Inches(4.95)).table
    col_widths = [1.15, 2.75, 2.6, 3.55, 1.1]
    for i, w in enumerate(col_widths):
        table.columns[i].width = Inches(w)

    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLORS["bg_2"]
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(10)
        p.font.name = "Aptos"
        p.font.color.rgb = COLORS["text_light"]
        p.alignment = PP_ALIGN.CENTER

    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(255, 255, 255) if r_idx % 2 else RGBColor(246, 248, 252)
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(9)
            p.font.name = "Aptos"
            p.font.color.rgb = COLORS["text_dark"]
            if c_idx == 4:
                p.font.bold = True
                p.alignment = PP_ALIGN.CENTER
                cell.fill.fore_color.rgb = RGBColor(231, 248, 237) if value == "PASS" else RGBColor(252, 233, 234)

    add_footer(slide, "5 PASS and 5 FAIL cases tailored to the current project state")


def challenges_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, dark=False)
    add_logo(slide)
    add_title(slide, "Challenges, Observations And Improvements", "Balanced review for viva discussion")

    add_card(slide, 0.72, 1.7, 5.8, 4.9, "Strengths", accent=COLORS["success"])
    add_bullets(
        slide,
        [
            "Strong functional coverage across operations, billing, analytics and reporting.",
            "Good desktop integration for startup management, PDF saving and backup handling.",
            "Clear separation of client, server and Electron layers.",
            "Data model is broad enough to support multiple customer/business types.",
        ],
        1.02,
        2.14,
        5.0,
        3.7,
        font_size=14,
    )

    add_card(slide, 6.75, 1.7, 5.82, 4.9, "Current Gaps / Next Improvements", accent=COLORS["accent_2"])
    add_bullets(
        slide,
        [
            "Add automated unit and integration tests for critical flows.",
            "Move fallback secrets and admin passwords out of source defaults.",
            "Re-enable or complete the license workflow if it is part of final business scope.",
            "Break the very large Admin.tsx file into smaller feature hooks/services for maintainability.",
            "Add audit logging and stronger validation around high-risk admin actions.",
        ],
        7.05,
        2.14,
        5.0,
        3.8,
        font_size=13.7,
    )
    add_footer(slide, "Useful slide for viva questions on limitations and software engineering maturity")


def conclusion_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, dark=True)
    add_logo(slide)

    add_paragraph(slide, "Conclusion", 0.82, 1.0, 2.8, 0.4, dark=True, size=26, bold=True)
    add_paragraph(
        slide,
        "Aachi Vanthachu is a well-scoped desktop ERP solution for magazine operations that combines customer onboarding, subscription billing, advertisement handling, inventory records, analytics, reporting and backup features in a single packaged system.",
        0.82,
        1.75,
        7.2,
        1.25,
        dark=True,
        size=15,
        color=RGBColor(217, 226, 239),
    )

    add_card(slide, 8.55, 1.35, 3.65, 2.0, dark=False, accent=COLORS["accent"])
    add_paragraph(slide, "Key Takeaway", 8.88, 1.7, 2.2, 0.25, size=15, bold=True)
    add_paragraph(
        slide,
        "The project demonstrates practical full-stack desktop engineering with real business workflows, not just a CRUD prototype.",
        8.88,
        2.05,
        2.75,
        0.65,
        size=12.5,
        color=COLORS["text_muted"],
    )

    add_metric(slide, 0.82, 4.35, 2.2, 1.05, "Integrated", "single-window business workflow", COLORS["accent"])
    add_metric(slide, 3.32, 4.35, 2.2, 1.05, "Automated", "PDF, Excel and backup tasks", COLORS["accent_2"])
    add_metric(slide, 5.82, 4.35, 2.2, 1.05, "Scalable", "modular collections and APIs", COLORS["accent_3"])

    add_paragraph(slide, "Thank You", 9.0, 5.15, 2.5, 0.35, dark=True, size=22, bold=True)
    add_paragraph(slide, "Questions and discussion", 9.0, 5.55, 2.9, 0.25, dark=True, size=13, color=RGBColor(201, 211, 225))
    add_footer(slide, "End of viva presentation", dark=True)


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    cover_slide(prs)
    agenda_slide(prs)
    stack_slide(prs)
    architecture_slide(prs)
    modules_slide(prs)
    workflow_slide(prs)
    data_model_slide(prs)
    analytics_slide(prs)
    security_slide(prs)
    testing_slide(prs)
    challenges_slide(prs)
    conclusion_slide(prs)

    prs.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    path = build_presentation()
    print(f"Presentation generated: {path}")
